#!/usr/bin/env python3
"""Acquire the official MHLW fatal occupational accident workbooks for 2019-2024.

This is an EVENTREG staging-corpus job. It preserves official source rows,
workbook/sheet lineage, schema drift, exact-duplicate diagnostics and cautious
Target48 discovery pointers. It never creates canonical events, 100 m cell
relations, scores, rankings or automatic exclusions.
"""
from __future__ import annotations

import csv
import gzip
import hashlib
import json
import re
import shutil
import sys
import time
import urllib.request
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

OUT = Path("output_mhlw")
RAW = OUT / "raw"
FILES = {
    2019: "https://anzeninfo.mhlw.go.jp/anzen/sib_xls/sibou_db_r01.xlsx",
    2020: "https://anzeninfo.mhlw.go.jp/anzen/sib_xls/sibou_db_r02.xlsx",
    2021: "https://anzeninfo.mhlw.go.jp/anzen/sib_xls/sibou_db_r03.xlsx",
    2022: "https://anzeninfo.mhlw.go.jp/anzen/sib_xls/sibou_db_r04.xlsx",
    2023: "https://anzeninfo.mhlw.go.jp/anzen/sib_xls/sibou_db_r05.xlsx",
    2024: "https://anzeninfo.mhlw.go.jp/anzen/sib_xls/sibou_db_r06.xlsx",
}

TOKYO_WARDS = [
    "千代田区", "中央区", "港区", "新宿区", "文京区", "台東区", "墨田区", "江東区",
    "品川区", "目黒区", "大田区", "世田谷区", "渋谷区", "中野区", "杉並区", "豊島区",
    "北区", "荒川区", "板橋区", "練馬区", "足立区", "葛飾区", "江戸川区",
]
YOKOHAMA_WARDS = [
    "鶴見区", "神奈川区", "西区", "中区", "南区", "港南区", "保土ケ谷区", "旭区",
    "磯子区", "金沢区", "港北区", "緑区", "青葉区", "都筑区", "戸塚区", "栄区",
    "泉区", "瀬谷区",
]
KAWASAKI_WARDS = ["川崎区", "幸区", "中原区", "高津区", "宮前区", "多摩区", "麻生区"]
AMBIGUOUS_WARDS = {"中央区", "港区", "北区", "西区", "中区", "南区", "旭区", "緑区", "青葉区", "泉区"}
HEADER_HINTS = (
    "ID", "災害", "発生", "年月", "業種", "起因物", "事故", "事業場", "規模", "分類",
    "都道府県", "市区町村", "所在地", "場所", "死亡", "死者",
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def download(url: str, destination: Path, retries: int = 6) -> dict[str, Any]:
    destination.parent.mkdir(parents=True, exist_ok=True)
    headers = {
        "User-Agent": "Mozilla/5.0 (compatible; EVENTREG-MHLW-harvester/1.0; research)",
        "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
        "Accept-Language": "ja,en;q=0.8",
    }
    last_error = ""
    for attempt in range(1, retries + 1):
        try:
            request = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(request, timeout=180) as response, destination.open("wb") as output:
                shutil.copyfileobj(response, output, length=1024 * 1024)
                receipt = {
                    "status": int(getattr(response, "status", 200)),
                    "content_type": response.headers.get("Content-Type", ""),
                    "final_url": response.geturl(),
                    "bytes": destination.stat().st_size,
                    "attempts": attempt,
                }
            if destination.stat().st_size < 1_000:
                raise RuntimeError(f"download too small: {destination.stat().st_size} bytes")
            return receipt
        except Exception as exc:
            last_error = repr(exc)
            destination.unlink(missing_ok=True)
            time.sleep(min(30, 2 ** attempt))
    raise RuntimeError(f"download failed: {url}: {last_error}")


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value).replace("\ufeff", "")).strip()


def normalized_header(value: Any, index: int) -> str:
    text = clean(value).replace("\n", " ")
    return text or f"column_{index + 1:03d}"


def unique_headers(values: tuple[Any, ...]) -> list[str]:
    counts: Counter[str] = Counter()
    result: list[str] = []
    for index, value in enumerate(values):
        base = normalized_header(value, index)
        counts[base] += 1
        result.append(base if counts[base] == 1 else f"{base}__{counts[base]}")
    return result


def header_score(values: tuple[Any, ...]) -> tuple[int, int]:
    texts = [clean(value) for value in values]
    nonempty = sum(bool(value) for value in texts)
    hints = sum(any(hint in value for hint in HEADER_HINTS) for value in texts if value)
    return hints, nonempty


def detect_header_row(worksheet, scan_limit: int = 40) -> tuple[int, list[str]]:
    candidates: list[tuple[int, int, int, tuple[Any, ...]]] = []
    for row_number, values in enumerate(
        worksheet.iter_rows(min_row=1, max_row=min(scan_limit, worksheet.max_row), values_only=True),
        start=1,
    ):
        hints, nonempty = header_score(values)
        if nonempty:
            candidates.append((hints, nonempty, -row_number, values))
    if not candidates:
        raise RuntimeError(f"no nonempty header candidate in sheet {worksheet.title}")
    _, _, negative_row, values = max(candidates)
    row_number = -negative_row
    headers = unique_headers(values)
    while headers and headers[-1].startswith("column_"):
        headers.pop()
    if not headers:
        raise RuntimeError(f"empty normalized header in sheet {worksheet.title}")
    return row_number, headers


def evidence_excerpt(text: str, token: str, width: int = 45) -> str:
    position = text.find(token)
    if position < 0:
        return ""
    return text[max(0, position - width): min(len(text), position + len(token) + width)]


def classify_target(fields: dict[str, str]) -> dict[str, str]:
    field_blob = " | ".join(f"{key}={value}" for key, value in fields.items() if value)
    compact = re.sub(r"\s+", "", field_blob)
    strong: list[tuple[str, str, str]] = []
    coarse: list[tuple[str, str, str]] = []
    weak: list[tuple[str, str, str]] = []

    for ward in TOKYO_WARDS:
        for token in (f"東京都{ward}", f"東京{ward}"):
            if token in compact:
                strong.append(("東京23区", ward, evidence_excerpt(compact, token)))
                break
    for ward in YOKOHAMA_WARDS:
        token = f"横浜市{ward}"
        if token in compact:
            strong.append(("横浜市", ward, evidence_excerpt(compact, token)))
    for ward in KAWASAKI_WARDS:
        token = f"川崎市{ward}"
        if token in compact:
            strong.append(("川崎市", ward, evidence_excerpt(compact, token)))

    if "東京都" in compact and not any(region == "東京23区" for region, _, _ in strong):
        coarse.append(("東京23区", "東京都（区未確定）", evidence_excerpt(compact, "東京都")))
    if "横浜市" in compact and not any(region == "横浜市" for region, _, _ in strong):
        coarse.append(("横浜市", "横浜市（区未確定）", evidence_excerpt(compact, "横浜市")))
    if "川崎市" in compact and not any(region == "川崎市" for region, _, _ in strong):
        coarse.append(("川崎市", "川崎市（区未確定）", evidence_excerpt(compact, "川崎市")))

    used = {municipality for _, municipality, _ in strong}
    for region, wards in (("東京23区", TOKYO_WARDS), ("横浜市", YOKOHAMA_WARDS), ("川崎市", KAWASAKI_WARDS)):
        for ward in wards:
            if ward in compact and ward not in used:
                weak.append((region, ward, evidence_excerpt(compact, ward)))

    def unique(items: list[tuple[str, str, str]]) -> list[tuple[str, str, str]]:
        seen: set[tuple[str, str]] = set()
        output = []
        for item in items:
            key = item[:2]
            if key not in seen:
                seen.add(key)
                output.append(item)
        return output

    strong, coarse, weak = unique(strong), unique(coarse), unique(weak)
    selected = strong or coarse or weak
    if not selected:
        return {
            "target_candidate": "0",
            "target_match_class": "NO_TARGET48_STRING",
            "target_region_group": "",
            "target_municipality_candidates": "",
            "target_match_evidence": "",
            "event_registry_eligibility": "SOURCE_CORPUS_ONLY",
            "geometry_eligibility": "NO",
            "privacy_gate": "ACTIVE",
            "unknown_reason": "no explicit target48 place string in official workbook row",
        }

    if strong:
        match_class = "EXPLICIT_CITY_PREFECTURE_PLUS_WARD"
        eligibility = "DISCOVERY_REVIEW__IDENTITY_LOCATION_AND_DEDUP_REQUIRED"
    elif coarse:
        match_class = "EXPLICIT_CITY_OR_PREFECTURE_ONLY"
        eligibility = "HOLD_SUBMUNICIPAL_LOCATION"
    else:
        names = {municipality for _, municipality, _ in weak}
        match_class = "AMBIGUOUS_WARD_ONLY" if names & AMBIGUOUS_WARDS else "WARD_ONLY_WITHOUT_CITY_CONTEXT"
        eligibility = "HOLD_CONTEXT_DISAMBIGUATION"

    return {
        "target_candidate": "1",
        "target_match_class": match_class,
        "target_region_group": "|".join(sorted({region for region, _, _ in selected})),
        "target_municipality_candidates": "|".join(municipality for _, municipality, _ in selected),
        "target_match_evidence": " || ".join(evidence for _, _, evidence in selected if evidence),
        "event_registry_eligibility": eligibility,
        "geometry_eligibility": "NO_AUTOMATIC_100M_PROMOTION",
        "privacy_gate": "ACTIVE",
        "unknown_reason": "event identity, exact occurrence place, dedup and geometry remain pending",
    }


def exact_fingerprint(fields: dict[str, str]) -> str:
    payload = json.dumps(fields, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str], compressed: bool = False) -> None:
    if compressed:
        handle = gzip.open(path, "wt", encoding="utf-8", newline="", compresslevel=9)
    else:
        handle = path.open("w", encoding="utf-8", newline="")
    with handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fields})


def main() -> int:
    OUT.mkdir(exist_ok=True)
    RAW.mkdir(parents=True, exist_ok=True)

    all_rows: list[dict[str, Any]] = []
    candidates: list[dict[str, Any]] = []
    strong_candidates: list[dict[str, Any]] = []
    manifest: list[dict[str, Any]] = []
    schemas: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    dynamic_headers: list[str] = []
    dynamic_seen: set[str] = set()
    year_counts: Counter[int] = Counter()
    sheet_counts: Counter[str] = Counter()
    match_classes: Counter[str] = Counter()
    fingerprint_first: dict[str, str] = {}
    duplicate_rows: list[dict[str, Any]] = []

    for year, url in FILES.items():
        filename = f"sibou_db_{year}.xlsx"
        path = RAW / filename
        print(f"FETCH {year} {url}", flush=True)
        try:
            receipt = download(url, path)
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            errors.append({"year": year, "url": url, "error": repr(exc)})
            print(f"ERROR {year}: {exc!r}", file=sys.stderr, flush=True)
            continue

        file_row_count = 0
        sheet_receipts = []
        for worksheet in workbook.worksheets:
            try:
                header_row, headers = detect_header_row(worksheet)
            except Exception as exc:
                sheet_receipts.append({"sheet": worksheet.title, "status": "SKIPPED_NO_HEADER", "error": repr(exc)})
                continue

            for header in headers:
                if header not in dynamic_seen:
                    dynamic_seen.add(header)
                    dynamic_headers.append(header)

            sheet_row_count = 0
            for row_number, values in enumerate(
                worksheet.iter_rows(min_row=header_row + 1, values_only=True),
                start=header_row + 1,
            ):
                trimmed_values = list(values[: len(headers)])
                normalized = {header: clean(value) for header, value in zip(headers, trimmed_values)}
                if not any(normalized.values()):
                    continue
                nonempty_count = sum(bool(value) for value in normalized.values())
                if nonempty_count < 2:
                    continue

                fingerprint = exact_fingerprint(normalized)
                source_identity = f"MHLW_SIBOU_{year}_{worksheet.title}_{row_number}"
                duplicate_of = fingerprint_first.get(fingerprint, "")
                if not duplicate_of:
                    fingerprint_first[fingerprint] = source_identity
                else:
                    duplicate_rows.append({
                        "source_identity": source_identity,
                        "duplicate_of": duplicate_of,
                        "fingerprint": fingerprint,
                        "source_year": year,
                        "source_sheet": worksheet.title,
                        "source_row": row_number,
                    })

                target = classify_target(normalized)
                record: dict[str, Any] = {
                    "source_dataset": "MHLW_FATAL_OCCUPATIONAL_DATABASE_2019_2024",
                    "source_file": filename,
                    "source_url": url,
                    "source_year": year,
                    "source_sheet": worksheet.title,
                    "source_header_row": header_row,
                    "source_row_number": row_number,
                    "canonical_source_identity": source_identity,
                    "source_record_fingerprint": fingerprint,
                    "exact_duplicate_of": duplicate_of,
                    "exact_duplicate_status": "EXACT_DUPLICATE_SOURCE_ROW" if duplicate_of else "UNIQUE_SOURCE_ROW",
                }
                record.update(normalized)
                record.update(target)
                all_rows.append(record)
                if target["target_candidate"] == "1":
                    candidates.append(record)
                    match_classes[target["target_match_class"]] += 1
                    if target["target_match_class"] == "EXPLICIT_CITY_PREFECTURE_PLUS_WARD":
                        strong_candidates.append(record)
                sheet_row_count += 1
                file_row_count += 1

            sheet_counts[f"{year}:{worksheet.title}"] = sheet_row_count
            sheet_receipts.append({
                "sheet": worksheet.title,
                "status": "PARSED",
                "header_row": header_row,
                "headers": headers,
                "rows": sheet_row_count,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
            })
            schemas.append({
                "source_year": year,
                "source_file": filename,
                "sheet": worksheet.title,
                "header_row": header_row,
                "headers": headers,
                "parsed_rows": sheet_row_count,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
            })

        workbook.close()
        year_counts[year] = file_row_count
        manifest.append({
            "year": year,
            "source_file": filename,
            "source_url": url,
            "final_url": receipt["final_url"],
            "http_status": receipt["status"],
            "content_type": receipt["content_type"],
            "bytes": receipt["bytes"],
            "sha256": sha256(path),
            "parsed_rows": file_row_count,
            "sheets": json.dumps(sheet_receipts, ensure_ascii=False, separators=(",", ":")),
            "status": "ACQUIRED_AND_PARSED",
            "canonical_events_added": 0,
            "cell_relations_added": 0,
        })

    metadata_fields = [
        "source_dataset", "source_file", "source_url", "source_year", "source_sheet",
        "source_header_row", "source_row_number", "canonical_source_identity",
        "source_record_fingerprint", "exact_duplicate_of", "exact_duplicate_status",
    ]
    gate_fields = [
        "target_candidate", "target_match_class", "target_region_group",
        "target_municipality_candidates", "target_match_evidence",
        "event_registry_eligibility", "geometry_eligibility", "privacy_gate", "unknown_reason",
    ]
    fields = metadata_fields + dynamic_headers + gate_fields

    write_csv(OUT / "MHLW_FATAL_OCCUPATIONAL_2019_2024_ALL_SOURCE_ROWS.csv.gz", all_rows, fields, True)
    write_csv(OUT / "MHLW_FATAL_OCCUPATIONAL_TARGET48_DISCOVERY_CANDIDATES.csv", candidates, fields)
    write_csv(OUT / "MHLW_FATAL_OCCUPATIONAL_TARGET48_STRONG_STRING_MATCHES.csv", strong_candidates, fields)
    write_csv(
        OUT / "MHLW_FATAL_OCCUPATIONAL_EXACT_DUPLICATE_MAP.csv",
        duplicate_rows,
        ["source_identity", "duplicate_of", "fingerprint", "source_year", "source_sheet", "source_row"],
    )
    write_csv(
        OUT / "SOURCE_FILE_MANIFEST.csv",
        manifest,
        [
            "year", "source_file", "source_url", "final_url", "http_status", "content_type",
            "bytes", "sha256", "parsed_rows", "sheets", "status",
            "canonical_events_added", "cell_relations_added",
        ],
    )
    write_csv(
        OUT / "YEAR_COUNTS.csv",
        [{"year": year, "records": count} for year, count in sorted(year_counts.items())],
        ["year", "records"],
    )
    write_csv(
        OUT / "SHEET_COUNTS.csv",
        [{"year_sheet": key, "records": value} for key, value in sorted(sheet_counts.items())],
        ["year_sheet", "records"],
    )
    (OUT / "SCHEMA_MANIFEST.json").write_text(
        json.dumps(schemas, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    shutil.make_archive(str(OUT / "MHLW_RAW_2019_2024_XLSX"), "zip", RAW)

    summary = {
        "run_id": "eventreg-mhlw-fatal-occupational-2019-2024-20260819-v1",
        "official_source": "MHLW workplace safety fatal accident annual workbooks",
        "coverage": "2019-2024",
        "expected_files": 6,
        "acquired_files": len(manifest),
        "acquired_source_rows": len(all_rows),
        "exact_duplicate_rows": len(duplicate_rows),
        "unique_source_fingerprints": len(fingerprint_first),
        "target48_discovery_candidates": len(candidates),
        "target48_strong_string_matches": len(strong_candidates),
        "target_match_classes": dict(match_classes),
        "year_counts": dict(year_counts),
        "errors": errors,
        "complete": len(manifest) == 6 and len(all_rows) > 0 and not errors,
        "contracts": {
            "canonical_events_added": 0,
            "cell_relations_added": 0,
            "score_rank_exclusion_effect": 0,
            "privacy_gate": "ACTIVE",
            "no_automatic_100m_promotion": True,
            "unknown_is_not_absence": True,
        },
    }
    (OUT / "SUMMARY.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "README.md").write_text(
        "# EVENTREG MHLW Fatal Occupational Bulk Harvest 2019–2024\n\n"
        f"- Files acquired: **{len(manifest)}/6**\n"
        f"- Source rows acquired: **{len(all_rows):,}**\n"
        f"- Exact duplicate source rows: **{len(duplicate_rows):,}**\n"
        f"- Target48 discovery candidates: **{len(candidates):,}**\n"
        f"- Strong city/prefecture + ward string matches: **{len(strong_candidates):,}**\n"
        "- Canonical events added: **0**\n"
        "- 100 m cell relations added: **0**\n\n"
        "Rows remain source-corpus records until independent event identity, exact occurrence place, "
        "deduplication, privacy and geometry adjudication are complete.\n",
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2), flush=True)
    return 0 if summary["complete"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
